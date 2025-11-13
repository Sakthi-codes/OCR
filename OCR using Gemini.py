import google.generativeai as genai
from docx import Document
import os
import datetime
from openpyxl import Workbook # New import for Excel

API_KEY = "API key" 
genai.configure(api_key=API_KEY)
GEMINI_MODEL = "gemini-2.5-flash-preview-05-20"

image_directory_path = "C:/Users/sakth/ocr_images" # Ensure this path is correct and accessible


SUPPORTED_EXTENSIONS = ('.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.webp')

INPUT_PRICE_PER_MILLION_USD = 0.15
OUTPUT_PRICE_PER_MILLION_USD = 0.60 # Using non-thinking output price



USD_TO_INR_RATE = 83.33 # Updated based on a more recent check (rates can vary slightly)

def extract_text_from_image_with_gemini(image_path):
    try:
        with open(image_path, "rb") as image_file:
            image_data = image_file.read()

        _, ext = os.path.splitext(image_path)
        
        mime_type = {
            '.jpg': 'image/jpeg',
            '.jpeg': 'image/jpeg',
            '.png': 'image/png',
            '.bmp': 'image/bmp',
            '.tiff': 'image/tiff',
            '.webp': 'image/webp'
        }.get(ext.lower(), 'application/octet-stream') # Default to generic if unknown

        image_part = {
            'mime_type': mime_type,
            'data': image_data
        }

        prompt_content = [
            image_part,
            "Please extract all the text visible in this image. Do not add any commentary, just the extracted text."
        ]

        model = genai.GenerativeModel(GEMINI_MODEL)

        response = model.generate_content(
            contents=prompt_content,
        )
        
        input_tokens = getattr(response.usage_metadata, 'prompt_token_count', 0)
        output_tokens = getattr(response.usage_metadata, 'candidates_token_count', 0)
        
        return response.text, input_tokens, output_tokens
    except Exception as e:
        print(f"Error extracting text from '{os.path.basename(image_path)}' with Gemini: {e}")
        return None, 0, 0 # Return None for text and 0 for tokens on error

def save_text_as_docx(text, output_filepath):
    try:
        document = Document()
        document.add_paragraph(text)
        document.save(output_filepath)
        print(f"Text successfully saved to: {output_filepath}")
        return True
    except Exception as e:
        print(f"Error saving DOCX file to '{output_filepath}': {e}")
        return False

if __name__ == "__main__":
    if not os.path.isdir(image_directory_path):
        print(f"Error: Directory not found at '{image_directory_path}'. Please provide a valid folder path.")
    else:
        print(f"Scanning directory: {image_directory_path} for images...")
        
        processed_files_data = [] # To store data for the Excel report
        processed_count = 0

        for filename in os.listdir(image_directory_path):
            if filename.lower().endswith(SUPPORTED_EXTENSIONS):
                full_image_path = os.path.join(image_directory_path, filename)
                print(f"\n--- Processing image: {filename} ---")
                
                extracted_text, input_tokens, output_tokens = extract_text_from_image_with_gemini(full_image_path)
                
                file_status = "Failed"
                usd_cost = 0.0
                inr_cost = 0.0
                total_tokens = input_tokens + output_tokens

                if extracted_text:
                    print("Extracted Text Preview:")
                    print(extracted_text[:200] + "..." if len(extracted_text) > 200 else extracted_text)
                    
                    image_filename_without_ext = os.path.splitext(filename)[0]
                    output_docx_filename = f"{image_filename_without_ext}_extracted_text.docx"
                    output_docx_filepath = os.path.join(image_directory_path, output_docx_filename)
                    
                    if save_text_as_docx(extracted_text, output_docx_filepath):
                        file_status = "Success"
                        
                        usd_cost = (input_tokens / 1_000_000) * INPUT_PRICE_PER_MILLION_USD + \
                                   (output_tokens / 1_000_000) * OUTPUT_PRICE_PER_MILLION_USD
                        inr_cost = usd_cost * USD_TO_INR_RATE
                        
                        processed_count += 1
                else:
                    print(f"No text extracted or an error occurred for '{filename}'.")
                
                processed_files_data.append({
                    "File ID": filename,
                    "Status": file_status,
                    "Input Tokens": input_tokens,
                    "Output Tokens": output_tokens,
                    "Total Tokens": total_tokens,
                    "USD Cost": usd_cost, # Keep as float for Excel
                    "INR Cost": inr_cost  # Keep as float for Excel
                })
            else:
                print(f"Skipping non-image file: {filename}")
        
        print(f"\n--- Processing Complete ---")
        print(f"Successfully processed {processed_count} images in '{image_directory_path}'.")

        if processed_files_data:
            report_filename = f"ocr_cost_report_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            report_filepath = os.path.join(image_directory_path, report_filename)
            
            excel_columns = ["File ID", "Status", "Input Tokens", "Output Tokens", "Total Tokens", "USD Cost", "INR Cost"]
            
            try:
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "OCR Cost Analysis" # Set sheet title

                sheet.append(excel_columns)

                for row_data in processed_files_data:
                    sheet.append([row_data[col] for col in excel_columns])
                
                workbook.save(report_filepath)
                print(f"Cost analysis report saved to: {report_filepath}")
            except Exception as e:
                print(f"Error writing Excel file: {e}")
        else:
            print("No image files processed to generate a report.")